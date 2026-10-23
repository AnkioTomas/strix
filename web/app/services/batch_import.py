"""Parse the batch-scan spreadsheet (应用名称 / 端口 / BU / 负责人 / URL)."""

from __future__ import annotations

import io
import re
import sys
from dataclasses import dataclass
from urllib.parse import urlsplit

_ALIASES = {
    "name": ("应用名称", "应用名", "应用", "系统名称", "名称", "name"),
    "port": ("端口", "port"),
    "bu": ("bu", "事业部", "业务单元"),
    "owner": ("负责人", "责任人", "owner", "邮箱", "负责人邮箱"),
    "url": ("url", "地址", "链接", "访问地址", "域名"),
}
_URL_HINTS = ("url", "地址", "链接", "域名", "网址")
_URL_LIKE = re.compile(r"^(https?://\S+|[\w.-]+\.[a-z]{2,}(:\d+)?(/\S*)?|\d{1,3}(\.\d{1,3}){3}(:\d+)?(/\S*)?)$", re.I)
_POSITIONAL = ("name", "port", "bu", "owner", "url")
_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class SheetError(ValueError):
    pass


@dataclass
class BatchRow:
    row: int
    name: str
    target: str
    bu: str
    owner: str
    emails: list[str]


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def build_target(url: str, port: str) -> str:
    """URL/host plus the port column; a port already in the URL wins."""
    url = url.strip()
    port = port.strip()
    if not url or not port.isdigit():
        return url
    parts = urlsplit(url if "://" in url else f"//{url}")
    if parts.port is not None or not parts.hostname:
        return url
    host = f"[{parts.hostname}]" if ":" in parts.hostname else parts.hostname
    netloc = f"{host}:{port}"
    scheme = f"{parts.scheme}://" if parts.scheme else ""
    rest = parts.path + (f"?{parts.query}" if parts.query else "")
    return f"{scheme}{netloc}{rest}"


def owner_emails(owner: str, mail_domain: str = "") -> list[str]:
    emails: list[str] = []
    for part in re.split(r"[,;，；\s]+", owner):
        part = part.strip()
        if not part:
            continue
        if "@" not in part and mail_domain:
            part = f"{part}@{mail_domain.lstrip('@')}"
        if _EMAIL.match(part):
            emails.append(part)
    return emails


def _load_openpyxl():
    """openpyxl optionally imports numpy; a numpy that crashes on this CPU
    (RuntimeError, not ImportError) must not break Excel import."""
    try:
        from openpyxl import load_workbook
    except RuntimeError:
        for name in [m for m in sys.modules if m == "numpy" or m.startswith(("numpy.", "openpyxl"))]:
            del sys.modules[name]
        sys.modules["numpy"] = None  # type: ignore[assignment]
        from openpyxl import load_workbook
    return load_workbook


def _find_url(row: list[str], exclude: set[int]) -> str:
    """URL column empty: use a cell outside the mapped columns that looks like one."""
    for idx, cell in enumerate(row):
        if idx not in exclude and _URL_LIKE.match(cell.strip()):
            return cell.strip()
    return ""


def parse_sheet(data: bytes, *, mail_domain: str = "") -> list[BatchRow]:
    load_workbook = _load_openpyxl()

    try:
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    except Exception as exc:
        raise SheetError("无法读取 Excel 文件（需要 .xlsx）") from exc
    rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    rows = [(i, r) for i, r in enumerate(rows, start=1) if any(r)]
    if not rows:
        raise SheetError("表格为空")

    header: dict[str, int] = {}
    first = [c.strip().lower() for c in rows[0][1]]
    for key, aliases in _ALIASES.items():
        for idx, title in enumerate(first):
            if idx in header.values():
                continue
            if title in aliases:
                header[key] = idx
                break
    if "url" not in header:
        # Tolerate titles like "应用URL地址" / "访问域名".
        for idx, title in enumerate(first):
            if idx not in header.values() and any(k in title for k in _URL_HINTS):
                header["url"] = idx
                break
    if header:
        body = rows[1:]
    else:
        # No recognisable header: fall back to column order A..E.
        header = {key: idx for idx, key in enumerate(_POSITIONAL)}
        body = rows

    def pick(row: list[str], key: str) -> str:
        idx = header.get(key)
        return row[idx] if idx is not None and idx < len(row) else ""

    result: list[BatchRow] = []
    for number, row in body:
        url = pick(row, "url") or _find_url(row, exclude=set(header.values()))
        owner = pick(row, "owner")
        if not url and not pick(row, "name"):
            # Stray cells (notes, status columns) on an otherwise empty row.
            continue
        result.append(
            BatchRow(
                row=number,
                name=pick(row, "name"),
                target=build_target(url, pick(row, "port")),
                bu=pick(row, "bu"),
                owner=owner,
                emails=owner_emails(owner, mail_domain),
            )
        )
    if not result:
        raise SheetError("表格没有数据行")
    return result
